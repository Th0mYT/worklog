import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

from summarizer.timesheet_export import export_month, match_rows, parse_hours, parse_summary, round_grain

ROWS = ['DDH - Libertas', 'DDH - K•HERO', 'DDH - K•HERO SyDiDOa', 'DDH - NeoNat',
        'DDH - Alvamiva - Agenas', 'Drop - DEM', 'Contractors']


class Parsing(unittest.TestCase):
    def test_durations_in_every_format_seen_in_the_sheets(self):
        for text, want in [('2h', 2), ('0.5h', 0.5), ('3,75h', 3.75), ('1h 55m', 1 + 55 / 60),
                           ('25m', 25 / 60), ('1,5', 1.5), ('2h 15m', 2.25)]:
            self.assertAlmostEqual(parse_hours(text), want, places=3, msg=text)
        self.assertIsNone(parse_hours('KH-1413'))
        self.assertIsNone(parse_hours('x2'))

    def test_sections_hours_and_total_line(self):
        text = ('## Progetto K-HERO\nDev work (1.5h) [khero-backend]\nReview (0.25h)\n'
                '## Generale\nMail (0.5h)\n\nTotal: 2.25h\n')
        a, b = parse_summary(text)
        self.assertEqual((a.header, a.hours, len(a.lines)), ('Progetto K-HERO', 1.75, 2))
        self.assertEqual((b.header, b.hours), ('Generale', 0.5))

    def test_a_parenthesis_inside_the_text_is_not_a_duration(self):
        s = parse_summary('## X\nFix (KH-1) in auth (2h) [a]\n')[0]
        self.assertEqual(s.hours, 2)

    def test_rounding_is_per_total_not_per_line(self):
        self.assertEqual(round_grain(0.1 + 0.1 + 0.1 + 0.1), 0.5)
        self.assertEqual(round_grain(0), 0)


class Matching(unittest.TestCase):
    def test_commesse_find_their_rows(self):
        got = match_rows(['Progetto K-HERO', 'Almaviva - Progetto Agenas - L1 implementation - Tech Lead',
                          'Totally unrelated'], ROWS)
        self.assertEqual(got, {'Progetto K-HERO': 'DDH - K•HERO',
                               'Almaviva - Progetto Agenas - L1 implementation - Tech Lead': 'DDH - Alvamiva - Agenas'})

    def test_explicit_mapping_wins(self):
        got = match_rows(['Progetto K-HERO'], ROWS, {'Progetto K-HERO': 'DDH - K•HERO SyDiDOa'})
        self.assertEqual(got['Progetto K-HERO'], 'DDH - K•HERO SyDiDOa')


class Export(unittest.TestCase):
    def build_template(self, path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A7'] = 'PROJECT'
        for i, label in enumerate(['DDH - K•HERO', 'DDH - Alvamiva - Agenas'], start=8):
            ws.cell(i, 1).value = label
        ws['AG8'] = '=IFERROR(SUM(B8:AF8), "")'        # AG9 deliberately missing
        ws['A11'] = 'DAILY TOTAL'
        ws['A14'] = 'PROJECT'
        ws['A15'], ws['A16'] = 'DDH - K•HERO', 'DDH - Almaviva - Agenas'   # label spelled differently
        ws['B8'] = 99                                                       # stale value must go
        wb.save(path)

    def test_fills_hours_notes_formulas_and_pending_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            self.build_template(tmp / 't.xlsx')
            (tmp / 's').mkdir()
            (tmp / 's' / '2026-10-02.md').write_text(
                '## Progetto K-HERO\nDev (1.5h) [khero-backend]\nReview (0.25h)\n'
                '## Almaviva - Progetto Agenas - L1 implementation - Tech Lead\nCall (0.5h)\n'
                '## Generale\nMisc (0.25h)\n\nTotal: 2.5h\n')
            report = export_month(2026, 10, tmp / 't.xlsx', tmp / 'out.xlsx', summaries_dir=tmp / 's')
            ws = openpyxl.load_workbook(tmp / 'out.xlsx').active
            self.assertEqual(ws['C8'].value, 1.75)
            self.assertIsNone(ws['B8'].value)
            self.assertEqual(ws['C9'].value, 0.5)
            self.assertEqual(ws['C15'].value, 'Dev (1.5h) [khero-backend]\nReview (0.25h)')
            self.assertEqual(ws['C16'].value, 'Call (0.5h)')
            self.assertEqual(ws['AG9'].value, '=IFERROR(SUM(B9:AF9), "")')
            self.assertEqual(report.formulas_added, ['DDH - Alvamiva - Agenas'])
            pending = openpyxl.load_workbook(tmp / 'out.xlsx')['Da assegnare']
            self.assertEqual([r[1] for r in pending.iter_rows(min_row=2, values_only=True)], ['Generale'])
            self.assertEqual(ws['G3'].value, 'October')

    def test_template_is_left_untouched(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            self.build_template(tmp / 't.xlsx')
            before = (tmp / 't.xlsx').read_bytes()
            (tmp / 's').mkdir()
            export_month(2026, 10, tmp / 't.xlsx', tmp / 'out.xlsx', summaries_dir=tmp / 's')
            self.assertEqual((tmp / 't.xlsx').read_bytes(), before)

    def test_not_a_timesheet_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            openpyxl.Workbook().save(tmp / 'x.xlsx')
            with self.assertRaises(ValueError):
                export_month(2026, 10, tmp / 'x.xlsx', tmp / 'o.xlsx', summaries_dir=tmp)


if __name__ == '__main__':
    unittest.main()


class Text(unittest.TestCase):
    LINES = ['Sviluppo backend (1,5h) [khero-backend]', 'Review PR (0.25h) [khero-backend, khero-angular]',
             'Review PR (0.25h) [khero-backend]', 'Call allineamento (1h 30m)', 'Appunti senza durata']

    def test_lines_normalises_durations_and_merges_repeats(self):
        from summarizer.timesheet_export import format_lines
        out = format_lines(self.LINES).split('\n')
        self.assertEqual(out, ['Sviluppo backend (1.5h) [khero-backend]',
                               'Review PR (0.5h) [khero-backend, khero-angular]',
                               'Call allineamento (1.5h)', 'Appunti senza durata'])

    def test_day_style_is_one_line_with_the_cell_total(self):
        from summarizer.timesheet_export import format_lines
        out = format_lines(self.LINES[:3], style='day', total=2.25)
        self.assertEqual(out, 'Sviluppo backend; Review PR (2.25h) [khero-backend, khero-angular]')

    def test_refs_can_be_dropped(self):
        from summarizer.timesheet_export import format_lines
        self.assertEqual(format_lines(['A (1h) [x]'], refs=False), 'A (1h)')

    def test_trailing_period_is_not_kept(self):
        from summarizer.timesheet_export import format_lines
        self.assertEqual(format_lines(['Fatto. (2.75h) [a]']), 'Fatto (2.75h) [a]')
