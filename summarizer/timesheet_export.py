"""
Month export — fills a copy of the shared timesheet workbook from the saved
daily summaries, so nothing has to be copied by hand.

  * the template is the user's own timesheet (any month): it is never modified,
    the result is always written to a new file
  * hours per day and commessa are the sum of the lines of that day's saved
    summary, rounded to the timesheet grain (0.25h) per commessa and day — not
    per block, so many short blocks no longer inflate the total
  * the matching row is found by name (see `match_rows`); the cell note holds the
    summary lines, in the "description (Xh) [refs]" format the sheet already uses
  * what has no row — unassigned work ("Generale"), commesse with no matching
    project row — is never guessed into the grid: it goes to a second sheet,
    "Da assegnare", with date, hours and text, for the user to place
  * every project row gets its TOTAL formula, so a row added by hand can no
    longer drop out of the monthly total

Usage:
  python -m summarizer.timesheet_export --month 2026-10 --template "<timesheet>.xlsx"
  python -m summarizer.timesheet_export --month 2026-10 --map "Progetto K-HERO=DDH - K•HERO"
  python -m summarizer.timesheet_export --month 2026-10 --style day --no-refs
"""

import calendar
import re
import sys
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from config import Config

GRAIN = 0.25
UNASSIGNED_HEADERS = {'generale', 'fuori commessa', 'non assegnata', 'general'}
_FIRST_COL, _LAST_COL, _TOTAL_COL = 2, 32, 33   # B..AF are days 1..31, AG is the row total
_MONTHS_IT = ['Gennaio', 'Febbraio', 'Marzo', 'Aprile', 'Maggio', 'Giugno', 'Luglio', 'Agosto',
              'Settembre', 'Ottobre', 'Novembre', 'Dicembre']

# ---------------------------------------------------------------------------
# Reading a saved summary
# ---------------------------------------------------------------------------

_DURATION = re.compile(r'\((?P<v>[^()]*?)\)(?=\s*(?:\[[^\]]*\])?\s*$)')
_H = re.compile(r'(\d+(?:[.,]\d+)?)\s*h', re.IGNORECASE)
_M = re.compile(r'(\d+)\s*m(?:in)?\b', re.IGNORECASE)


def parse_hours(text: str) -> float | None:
    """'2h', '0.5h', '3,75h', '1h 55m', '25m', '1,5' → hours; None if it is not a duration."""
    t = (text or '').strip()
    if re.fullmatch(r'\d+(?:[.,]\d+)?', t):          # "(1,5)" — the unit was left out
        return float(t.replace(',', '.'))
    h, m = _H.search(t), _M.search(t)
    if not (h or m):
        return None
    if re.sub(r'[\d.,\s]|h|min|m', '', t, flags=re.IGNORECASE):
        return None
    return (float(h.group(1).replace(',', '.')) if h else 0.0) + (int(m.group(1)) / 60 if m else 0.0)


@dataclass
class Section:
    header: str | None
    hours: float = 0.0
    lines: list[str] = field(default_factory=list)


def parse_summary(text: str) -> list[Section]:
    """Split a saved summary into sections (one per '## header'); the 'Total:' line is ignored."""
    sections: OrderedDict[str | None, Section] = OrderedDict()
    current: Section | None = None
    for raw in (text or '').splitlines():
        line = raw.strip()
        if not line or line.lower().startswith('total'):
            continue
        if line.startswith('#'):
            header = line.lstrip('#').strip() or None
            current = sections.setdefault(header, Section(header))
            continue
        if current is None:
            current = sections.setdefault(None, Section(None))
        current.lines.append(line)
        m = _DURATION.search(line)
        h = parse_hours(m.group('v')) if m else None
        if h is not None:
            current.hours += h
    return list(sections.values())


_LINE = re.compile(r'^(?P<text>.*?)\s*\((?P<dur>[^()]*)\)\s*(?P<refs>\[[^\]]*\])?\s*$')


def _fmt_h(h: float) -> str:
    return f'{round(h, 2):g}h'


def format_lines(lines: list[str], *, style: str = 'lines', refs: bool = True,
                 total: float | None = None) -> str:
    """Cell text for one commessa and day.

    'lines' keeps one line per activity, durations normalised to `1.5h` (the sheet
    mixes `1,5h`, `(1h 55m)` and a bare `(1,5)`) and identical lines merged.
    'day' joins everything into a single line `a; b; c (Xh) [refs]` — no model, the
    descriptions are only concatenated. `refs=False` drops the `[repo, app]` part.
    """
    items: list[tuple[str, float | None, list[str]]] = []
    for line in lines:
        m = _LINE.match(line)
        hours = parse_hours(m.group('dur')) if m else None
        if m and hours is not None:
            text = m.group('text').strip().rstrip('.;')
            tags = [t.strip() for t in (m.group('refs') or '')[1:-1].split(',') if t.strip()]
        else:
            text, hours, tags = line.strip(), None, []
        items.append((text, hours, tags))

    merged: OrderedDict[str, tuple[float | None, list[str]]] = OrderedDict()
    for text, hours, tags in items:
        if text in merged:
            old_h, old_tags = merged[text]
            h = (old_h or 0) + hours if hours is not None else old_h
            merged[text] = (h, old_tags + [t for t in tags if t not in old_tags])
        else:
            merged[text] = (hours, list(tags))

    def refs_part(tags: list[str]) -> str:
        return f' [{", ".join(tags)}]' if refs and tags else ''

    if style == 'day':
        all_tags: list[str] = []
        for _, tags in merged.values():
            all_tags += [t for t in tags if t not in all_tags]
        hours = total if total is not None else sum(h or 0 for h, _ in merged.values())
        text = '; '.join(merged)
        return f'{text} ({_fmt_h(hours)}){refs_part(all_tags)}' if hours else text + refs_part(all_tags)

    out = []
    for text, (hours, tags) in merged.items():
        out.append(f'{text} ({_fmt_h(hours)})' + refs_part(tags) if hours is not None else text)
    return '\n'.join(out)


def round_grain(hours: float) -> float:
    return round(round(hours / GRAIN) * GRAIN, 2) if hours > 0 else 0.0


# ---------------------------------------------------------------------------
# Matching commesse to timesheet rows
# ---------------------------------------------------------------------------

def _tokens(name: str) -> set[str]:
    cleaned = re.sub(r'[^0-9a-zà-ÿ]+', ' ', re.sub(r'(?<=\w)[•\-·](?=\w)', '', name.lower()))
    return {t for t in cleaned.split() if t not in {'progetto', 'drop', 'ddh', 'tech', 'lead', 'l1'}}


def match_rows(commesse: list[str], row_labels: list[str], explicit: dict[str, str] | None = None) -> dict[str, str]:
    """{commessa: row label}. Explicit pairs win, then a unique best token overlap; a tie or no overlap is left out."""
    explicit = explicit or {}
    labels = {l.lower(): l for l in row_labels}
    result: dict[str, str] = {}
    for c in commesse:
        if c in explicit and explicit[c].lower() in labels:
            result[c] = labels[explicit[c].lower()]
            continue
        ct = _tokens(c)
        scored = []
        for label in row_labels:
            lt = _tokens(label)
            inter = len(ct & lt)
            if inter:
                scored.append((inter / len(ct | lt), label))
        scored.sort(reverse=True)
        if scored and (len(scored) == 1 or scored[0][0] > scored[1][0]):
            result[c] = scored[0][1]
    return result


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

@dataclass
class ExportReport:
    path: Path
    days_with_summary: int = 0
    missing_days: list[date] = field(default_factory=list)
    mapping: dict[str, str] = field(default_factory=dict)
    hours_by_row: dict[str, float] = field(default_factory=dict)
    pending: list[tuple[date, str, float]] = field(default_factory=list)   # (day, section, hours) left out of the grid
    formulas_added: list[str] = field(default_factory=list)

    def render(self) -> str:
        out = [f'Written:   {self.path}', f'Days:      {self.days_with_summary} with a summary']
        if self.missing_days:
            out.append('No summary: ' + ', '.join(d.strftime('%d') for d in self.missing_days if d.weekday() < 5))
        for c, row in self.mapping.items():
            out.append(f'Row:       {c}  →  {row}  (row total {self.hours_by_row.get(row, 0):g}h)')
        if self.formulas_added:
            out.append('TOTAL formula added on: ' + ', '.join(self.formulas_added))
        if self.pending:
            total = sum(h for _, _, h in self.pending)
            out.append(f'To assign: {total:g}h in {len(self.pending)} sections — see sheet "Da assegnare"')
        return '\n'.join(out)


def _locate(ws) -> tuple[int, int, int, int]:
    """(header row, daily-total row, notes header row, first notes row) from the column-A labels."""
    header = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'PROJECT']
    daily = next((r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'DAILY TOTAL'), 0)
    if len(header) < 2 or not daily:
        raise ValueError('Template does not look like the shared timesheet (PROJECT / DAILY TOTAL rows not found).')
    return header[0], daily, header[1], header[1] + 1


def export_month(year: int, month: int, template: Path, out: Path, *, summaries_dir: Path | None = None,
                 explicit: dict[str, str] | None = None, employee: str | None = None,
                 style: str = 'lines', refs: bool = True) -> ExportReport:
    import openpyxl
    from openpyxl.styles import Alignment

    summaries_dir = summaries_dir or Path(Config.SUMMARIES_DIR)
    wb = openpyxl.load_workbook(template)
    ws = wb.active
    h1, daily, h2, notes0 = _locate(ws)

    project_rows = {str(ws.cell(r, 1).value).strip(): r for r in range(h1 + 1, daily) if ws.cell(r, 1).value}
    note_rows = {}
    for r in range(notes0, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label and str(label).startswith('INSTRUCTIONS'):
            break
        if label:
            note_rows[str(label).strip()] = r

    # The notes block repeats the project labels, but not always verbatim (a typo fixed
    # in one place only, a subtotal row with no notes): exact label first, then by name.
    notes_row: dict[str, int] = {}
    fuzzy = match_rows([l for l in project_rows if l not in note_rows], list(note_rows))
    for label in project_rows:
        key = label if label in note_rows else fuzzy.get(label)
        if key:
            notes_row[label] = note_rows[key]

    # new month on a copy of another month's sheet: day headers are formulas on E3/G3
    ws['E3'] = year
    ws['G3'] = calendar.month_name[month]
    if employee:
        ws['B3'] = employee
    for r in project_rows.values():
        for c in range(_FIRST_COL, _LAST_COL + 1):
            ws.cell(r, c).value = None
    for r in note_rows.values():
        for c in range(_FIRST_COL, _LAST_COL + 1):
            ws.cell(r, c).value = None

    days = range(1, calendar.monthrange(year, month)[1] + 1)
    parsed: dict[date, list[Section]] = {}
    report = ExportReport(path=out)
    for d in days:
        day = date(year, month, d)
        f = summaries_dir / f'{day}.md'
        if f.exists():
            parsed[day] = parse_summary(f.read_text(encoding='utf-8'))
        else:
            report.missing_days.append(day)
    report.days_with_summary = len(parsed)

    names = [s.header for secs in parsed.values() for s in secs
             if s.header and s.header.lower() not in UNASSIGNED_HEADERS]
    report.mapping = match_rows(list(OrderedDict.fromkeys(names)), list(project_rows), explicit)

    for day, sections in parsed.items():
        col = _FIRST_COL + day.day - 1
        per_row: dict[str, Section] = {}
        for s in sections:
            row_label = report.mapping.get(s.header or '')
            if row_label is None:
                report.pending.append((day, s.header or '(senza intestazione)', round_grain(s.hours)))
                continue
            agg = per_row.setdefault(row_label, Section(row_label))
            agg.hours += s.hours
            agg.lines.extend(s.lines)
        for row_label, agg in per_row.items():
            hours = round_grain(agg.hours)
            ws.cell(project_rows[row_label], col).value = hours or None
            report.hours_by_row[row_label] = report.hours_by_row.get(row_label, 0) + hours
            if row_label in notes_row:
                cell = ws.cell(notes_row[row_label], col)
                cell.value = format_lines(agg.lines, style=style, refs=refs, total=hours)
                cell.alignment = Alignment(wrap_text=True, vertical='top')

    for label, r in project_rows.items():
        if ws.cell(r, _TOTAL_COL).value in (None, ''):
            ws.cell(r, _TOTAL_COL).value = f'=IFERROR(SUM(B{r}:AF{r}), "")'
            report.formulas_added.append(label)

    sheet = wb.create_sheet('Da assegnare')
    sheet.append(['Data', 'Sezione', 'Ore', 'Testo'])
    for day, sections in parsed.items():
        for s in sections:
            if report.mapping.get(s.header or '') is None:
                sheet.append([day.isoformat(), s.header or '', round_grain(s.hours),
                              format_lines(s.lines, style=style, refs=refs)])
    sheet.column_dimensions['D'].width = 120
    for row in sheet.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical='top')

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return report


def default_out(year: int, month: int) -> Path:
    base = Path(Config.SUMMARIES_DIR).parent / 'exports'
    path = base / f'{_MONTHS_IT[month - 1]}_{year}.xlsx'
    n = 2
    while path.exists():
        path = base / f'{_MONTHS_IT[month - 1]}_{year}_{n}.xlsx'
        n += 1
    return path


def main(argv: list[str] | None = None) -> int:
    import argparse
    p = argparse.ArgumentParser(description='Fill a copy of the shared timesheet from the saved summaries')
    p.add_argument('--month', required=True, metavar='YYYY-MM')
    p.add_argument('--template', type=Path, default=Path(Config.TIMESHEET_TEMPLATE) if Config.TIMESHEET_TEMPLATE else None,
                   help='Any month of the shared timesheet (default: timesheet_template in config.toml)')
    p.add_argument('--out', type=Path, default=None, help='Output file (default: ~/.worklog/exports/<Mese>_<anno>.xlsx)')
    p.add_argument('--map', action='append', default=[], metavar='COMMESSA=ROW',
                   help='Force a commessa onto a row label; repeatable')
    p.add_argument('--employee', default=None)
    p.add_argument('--style', choices=['lines', 'day'], default='lines',
                   help="'lines': one line per activity (default); 'day': one line per commessa and day")
    p.add_argument('--no-refs', action='store_true', help='Leave the [repo, app] references out of the notes')
    args = p.parse_args(argv)

    try:
        year, month = (int(x) for x in args.month.split('-'))
        date(year, month, 1)
    except ValueError:
        print('ERROR: --month must be YYYY-MM', file=sys.stderr)
        return 2
    if not args.template or not args.template.exists():
        print('ERROR: template not found — pass --template or set timesheet_template in config.toml', file=sys.stderr)
        return 2
    explicit = dict(m.split('=', 1) for m in args.map if '=' in m)
    out = args.out or default_out(year, month)
    if out.resolve() == args.template.resolve():
        print('ERROR: the output must be a new file, not the template', file=sys.stderr)
        return 2
    try:
        report = export_month(year, month, args.template, out, explicit=explicit, employee=args.employee,
                              style=args.style, refs=not args.no_refs)
    except (ValueError, OSError) as e:
        print(f'ERROR: {e}', file=sys.stderr)
        return 1
    print(report.render())
    return 0


if __name__ == '__main__':
    sys.exit(main())
